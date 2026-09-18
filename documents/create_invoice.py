from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).with_name("facture_site_web_interface_client.docx")
BLUE = "007BFF"
NAVY = "17212B"
LIGHT_BLUE = "EAF4FD"
LIGHT_GRAY = "D9E3EC"


def set_cell_shading(cell, fill):
    """Apply a solid background color to a table cell."""
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_border(cell, color=LIGHT_GRAY):
    """Apply subtle borders to every edge of a table cell."""
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    """Add internal padding so table content remains easy to read."""
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        margin = margins.find(qn("w:" + side))
        if margin is None:
            margin = OxmlElement("w:" + side)
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def write_cell(cell, text, bold=False, color=NAVY, size=9, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    """Write formatted text into a table cell."""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_label_value(paragraph, label, value):
    """Add a compact label and editable placeholder value."""
    label_run = paragraph.add_run(label)
    label_run.bold = True
    label_run.font.color.rgb = RGBColor.from_string(NAVY)
    value_run = paragraph.add_run(value)
    value_run.font.color.rgb = RGBColor.from_string("52616E")


document = Document()
section = document.sections[0]
section.top_margin = Cm(1.65)
section.bottom_margin = Cm(1.65)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(9.5)

# Create a clear invoice title and reference block.
title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(2)
run = title.add_run("Facture")
run.font.name = "Aptos Display"
run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
run.font.size = Pt(30)
run.font.bold = True
run.font.color.rgb = RGBColor.from_string(NAVY)

subtitle = document.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
subtitle_run = subtitle.add_run("Création d'un site web de quatre pages et d'une interface client")
subtitle_run.font.name = "Aptos"
subtitle_run.font.size = Pt(11)
subtitle_run.font.color.rgb = RGBColor.from_string("52616E")

metadata = document.add_table(rows=1, cols=2)
metadata.autofit = False
metadata.columns[0].width = Cm(8.1)
metadata.columns[1].width = Cm(8.1)
issuer = metadata.cell(0, 0)
client = metadata.cell(0, 1)
for cell in (issuer, client):
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_border(cell)

write_cell(issuer, "ÉMETTEUR", bold=True, color=BLUE, size=8)
for text in ("[Votre nom ou raison sociale]", "[Votre adresse]", "N° d'entreprise / TVA : [à compléter]", "E-mail : [à compléter] | Tél. : [à compléter]"):
    paragraph = issuer.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text).font.size = Pt(9)

write_cell(client, "CLIENT", bold=True, color=BLUE, size=8)
for text in ("[Nom du client ou société]", "[Adresse du client]", "N° TVA : [à compléter si applicable]", "E-mail : [à compléter]"):
    paragraph = client.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text).font.size = Pt(9)

document.add_paragraph().paragraph_format.space_after = Pt(0)

reference = document.add_table(rows=2, cols=4)
reference.autofit = False
reference.columns[0].width = Cm(3.0)
reference.columns[1].width = Cm(4.6)
reference.columns[2].width = Cm(3.0)
reference.columns[3].width = Cm(4.6)
reference_values = [
    ("N° facture", "FAC-2026-001", "Date", "18 septembre 2026"),
    ("Échéance", "18 octobre 2026", "Devise", "EUR"),
]
for row, values in zip(reference.rows, reference_values):
    for index, value in enumerate(values):
        cell = row.cells[index]
        set_cell_border(cell)
        if index % 2 == 0:
            set_cell_shading(cell, "F4F8FB")
            write_cell(cell, value, bold=True, color=NAVY, size=8.5)
        else:
            write_cell(cell, value, size=8.5)

document.add_paragraph().paragraph_format.space_after = Pt(0)

# List the billed services with a concise pricing table.
items = document.add_table(rows=1, cols=3)
items.autofit = False
items.columns[0].width = Cm(10.6)
items.columns[1].width = Cm(2.6)
items.columns[2].width = Cm(2.6)
headers = ("PRESTATION", "QTÉ", "MONTANT HT")
for cell, label in zip(items.rows[0].cells, headers):
    set_cell_shading(cell, NAVY)
    set_cell_border(cell, NAVY)
    write_cell(cell, label, bold=True, color="FFFFFF", size=8.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)

lines = [
    ("Création d'un site web vitrine de quatre pages\nConception visuelle, intégration responsive, navigation, formulaire de contact et SEO de base.", "1", "1 690,00 €"),
    ("Conception et intégration d'une interface client\nEspace de consultation dédié, parcours utilisateur et intégration des écrans prévus au périmètre du projet.", "1", "1 200,00 €"),
]
for index, row_values in enumerate(lines):
    row = items.add_row()
    fill = "FFFFFF" if index % 2 == 0 else "F6FAFD"
    for cell, value in zip(row.cells, row_values):
        set_cell_shading(cell, fill)
        set_cell_border(cell)
        alignment = WD_ALIGN_PARAGRAPH.LEFT if cell is row.cells[0] else WD_ALIGN_PARAGRAPH.CENTER
        write_cell(cell, value, size=9, alignment=alignment)

document.add_paragraph().paragraph_format.space_after = Pt(0)

totals = document.add_table(rows=3, cols=2)
totals.autofit = False
totals.columns[0].width = Cm(12.0)
totals.columns[1].width = Cm(3.8)
total_rows = [
    ("Total HT", "2 890,00 €", False),
    ("TVA", "[à compléter]", False),
    ("Total TTC", "[à compléter]", True),
]
for row, (label, value, is_total) in zip(totals.rows, total_rows):
    for cell in row.cells:
        set_cell_border(cell, NAVY if is_total else LIGHT_GRAY)
        set_cell_shading(cell, "EAF4FD" if is_total else "FFFFFF")
    write_cell(row.cells[0], label, bold=is_total, size=10, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    write_cell(row.cells[1], value, bold=True, size=10, alignment=WD_ALIGN_PARAGRAPH.RIGHT)

document.add_paragraph().paragraph_format.space_after = Pt(5)

payment = document.add_paragraph()
payment.paragraph_format.space_after = Pt(4)
payment_run = payment.add_run("Modalités de paiement")
payment_run.bold = True
payment_run.font.size = Pt(10.5)
payment_run.font.color.rgb = RGBColor.from_string(NAVY)

payment_text = document.add_paragraph()
payment_text.paragraph_format.space_after = Pt(2)
add_label_value(payment_text, "Règlement : ", "virement bancaire à 30 jours date de facture.")
bank_text = document.add_paragraph()
bank_text.paragraph_format.space_after = Pt(0)
add_label_value(bank_text, "IBAN : ", "[à compléter]     ")
add_label_value(bank_text, "BIC : ", "[à compléter]")

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_paragraph.add_run("Merci pour votre confiance")
footer_run.font.name = "Aptos"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor.from_string("71808C")

document.core_properties.title = "Facture Site Web et Interface Client"
document.core_properties.subject = "Facture pour la création d'un site web de quatre pages et d'une interface client"
document.save(OUTPUT)
print(OUTPUT)
